from __future__ import annotations
import io
import pytest
import pytest_asyncio
import openpyxl
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from sqlalchemy import select

from database.models import Base, UserStudyProfile, DailyStudySession, StudyTopicItem, StudyStreak, StudyBadge
from services.topic_service import topic_service
from services.export_service import export_service

TEST_DB_URL = "sqlite+aiosqlite:///:memory:"

@pytest_asyncio.fixture
async def db_session():
    engine = create_async_engine(TEST_DB_URL, echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with async_session() as session:
        yield session

    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
    await engine.dispose()

@pytest.mark.asyncio
async def test_edit_study_session(db_session: AsyncSession):
    """Test editing an existing study log adjusts topics, duration, problems, and notes."""
    discord_id = "user_edit_123"

    # 1. Create initial log
    log1 = await topic_service.log_study_activity(
        db=db_session,
        discord_id=discord_id,
        raw_topics="1. Factory Pattern 2. Singleton Pattern",
        notes="Initial notes on creational patterns",
        duration_minutes=40,
        problems_solved=1,
        category="LLD",
        username="dev_kunal",
        display_name="Kunal Dev"
    )

    session_id = log1["session_id"]
    assert log1["total_topics"] == 2
    assert log1["total_hours"] == 0.7

    # 2. Edit session: expand to 3 topics, 90 mins, 4 problems
    edit_res = await topic_service.edit_study_session(
        db=db_session,
        discord_id=discord_id,
        session_id=session_id,
        new_raw_topics="1. Factory Pattern\n2. Singleton Pattern\n3. Builder Pattern",
        new_notes="Updated notes with Builder pattern added",
        new_duration=90,
        new_problems=4,
        new_category="System Design"
    )

    assert edit_res["session_id"] == session_id
    assert len(edit_res["topics"]) == 3
    assert edit_res["duration_minutes"] == 90
    assert edit_res["problems_solved"] == 4
    assert edit_res["category"] == "System Design"
    assert edit_res["total_topics"] == 3
    assert edit_res["total_hours"] == 1.5

    # Check topic items in DB
    t_stmt = select(StudyTopicItem).where(StudyTopicItem.session_id == session_id)
    items = list((await db_session.execute(t_stmt)).scalars().all())
    assert len(items) == 3
    topic_names = [t.topic_name for t in items]
    assert "Builder Pattern" in topic_names

@pytest.mark.asyncio
async def test_delete_study_session(db_session: AsyncSession):
    """Test deleting an accidental study session cleanly adjusts cumulative metrics."""
    discord_id = "user_del_456"

    # Log two sessions
    log1 = await topic_service.log_study_activity(
        db=db_session,
        discord_id=discord_id,
        raw_topics="1. Arrays 2. Strings",
        notes="Basics",
        duration_minutes=30,
        problems_solved=2,
        category="DSA",
        username="alice",
        display_name="Alice"
    )

    log2 = await topic_service.log_study_activity(
        db=db_session,
        discord_id=discord_id,
        raw_topics="1. Trees 2. Graphs 3. Tries",
        notes="Advanced data structures",
        duration_minutes=60,
        problems_solved=5,
        category="DSA",
        username="alice",
        display_name="Alice"
    )

    # Total should be 5 topics, 90 minutes, 7 problems
    prof_stmt = select(UserStudyProfile).where(UserStudyProfile.discord_id == discord_id)
    prof = (await db_session.execute(prof_stmt)).scalar_one()
    assert prof.total_topics_count == 5
    assert prof.total_study_minutes == 90
    assert prof.total_problems_solved == 7

    # Delete session 1
    del_res = await topic_service.delete_study_session(
        db=db_session,
        discord_id=discord_id,
        session_id=log1["session_id"]
    )

    assert del_res["remaining_total_topics"] == 3
    assert del_res["remaining_total_hours"] == 1.0
    assert del_res["deleted_duration"] == 30

    # Ensure session 1 is deleted from DB
    sess_stmt = select(DailyStudySession).where(DailyStudySession.id == log1["session_id"])
    assert (await db_session.execute(sess_stmt)).scalar_one_or_none() is None

    # Ensure topic items for session 1 are deleted
    t_stmt = select(StudyTopicItem).where(StudyTopicItem.session_id == log1["session_id"])
    assert len(list((await db_session.execute(t_stmt)).scalars().all())) == 0

@pytest.mark.asyncio
async def test_reset_user_progress(db_session: AsyncSession):
    """Test full progress reset clears all sessions, topics, streak, and resets profile metrics."""
    discord_id = "user_reset_789"

    # Seed study activity
    await topic_service.log_study_activity(
        db=db_session,
        discord_id=discord_id,
        raw_topics="1. Kafka PubSub 2. RabbitMQ 3. Redis Streams",
        notes="Messaging architectures",
        duration_minutes=120,
        problems_solved=3,
        category="HLD",
        username="bob",
        display_name="Bob Engineer"
    )

    # Reset progress
    reset_ok = await topic_service.reset_user_progress(db=db_session, discord_id=discord_id)
    assert reset_ok is True

    # Check that sessions are gone
    s_stmt = select(DailyStudySession).where(DailyStudySession.discord_id == discord_id)
    assert len(list((await db_session.execute(s_stmt)).scalars().all())) == 0

    # Check that topics are gone
    t_stmt = select(StudyTopicItem).where(StudyTopicItem.discord_id == discord_id)
    assert len(list((await db_session.execute(t_stmt)).scalars().all())) == 0

    # Check streak is 0
    st_stmt = select(StudyStreak).where(StudyStreak.discord_id == discord_id)
    streak = (await db_session.execute(st_stmt)).scalar_one()
    assert streak.current_streak == 0
    assert streak.longest_streak == 0
    assert streak.total_days_studied == 0

    # Check profile stats
    p_stmt = select(UserStudyProfile).where(UserStudyProfile.discord_id == discord_id)
    prof = (await db_session.execute(p_stmt)).scalar_one()
    assert prof.total_topics_count == 0
    assert prof.total_study_minutes == 0
    assert prof.total_problems_solved == 0
    assert prof.rank_level == 0

@pytest.mark.asyncio
async def test_excel_export_generation(db_session: AsyncSession):
    """Test generating full Excel (.xlsx) workbook with openpyxl."""
    discord_id = "user_export_999"

    # Add 2 study sessions
    await topic_service.log_study_activity(
        db=db_session,
        discord_id=discord_id,
        raw_topics="1. Dynamic Programming 2. Backtracking: N-Queens",
        notes="Recursion trees and memoization tables",
        duration_minutes=60,
        problems_solved=2,
        category="DSA",
        username="charlie",
        display_name="Charlie Coder"
    )

    await topic_service.log_study_activity(
        db=db_session,
        discord_id=discord_id,
        raw_topics="1. Microservices Architecture: Saga Pattern",
        notes="Distributed transactions with choreo and orchestrator",
        duration_minutes=50,
        problems_solved=0,
        category="HLD",
        username="charlie",
        display_name="Charlie Coder"
    )

    # 1. Fetch export data
    export_data = await topic_service.get_full_export_data(db=db_session, discord_id=discord_id)
    assert len(export_data["sessions"]) == 2
    assert len(export_data["topics"]) == 3
    assert export_data["profile"]["total_topics_count"] == 3

    # 2. Generate Excel workbook stream
    excel_stream = export_service.generate_excel_export(export_data, display_name="Charlie Coder")
    assert isinstance(excel_stream, io.BytesIO)
    excel_bytes = excel_stream.getvalue()
    assert len(excel_bytes) > 0

    # 3. Read back Excel workbook using openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Overview & Scorecard" in sheet_names
    assert "Study Logs" in sheet_names
    assert "Topics & Notes Archive" in sheet_names

    # Check overview sheet content
    overview_sheet = wb["Overview & Scorecard"]
    assert "STUDY TRACKER & PREPARATION REPORT" in str(overview_sheet["A1"].value)

    # Check logs sheet content
    logs_sheet = wb["Study Logs"]
    assert logs_sheet.max_row >= 3 # Header + 2 rows of data

    # Check notes archive content
    notes_sheet = wb["Topics & Notes Archive"]
    assert notes_sheet.max_row >= 4 # Header + 3 topics
